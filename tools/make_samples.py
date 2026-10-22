# Örnek (kurgusal) veri dosyalarını üretir: samples/
# Kullanım: python3 tools/make_samples.py
# Gerçek öğrenci bilgisi içermez. Ders programı Excel'leri bölümün yayınladığı dosyaların kopyasıdır.
import json, shutil, datetime as dt, random
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'samples'
OUT.mkdir(exist_ok=True)
MUF = json.loads((ROOT / 'data/mufredat.json').read_text())
COURSES = {c['code']: c for c in MUF['courses']}
pdfmetrics.registerFont(TTFont('DV', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
pdfmetrics.registerFont(TTFont('DVB', '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))
HEAD = Font(bold=True, color='FFFFFF')
FILL = PatternFill('solid', fgColor='1E2B4D')
random.seed(7)


def sheet(wb, title, header, rows, widths=None, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.font = HEAD; c.fill = FILL; c.alignment = Alignment(vertical='center')
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    return ws


# ---------- 1. Müfredat ----------
def make_curriculum():
    wb = openpyxl.Workbook()
    sem_of = {}
    for s in MUF['semesters']:
        for c in s['mandatory']:
            sem_of[c] = s['semester']
    rows = []
    for c in sorted(MUF['courses'], key=lambda c: (sem_of.get(c['code'], 99), c['code'])):
        rows.append([sem_of.get(c['code'], ''), c['code'], c['name'], c.get('nameEn', ''), c.get('tul', ''), c.get('credit', 0), c['ects'], 'Zorunlu' if c['code'] in sem_of else 'Seçmeli'])
    sheet(wb, 'Dersler', ['Yarıyıl', 'Ders Kodu', 'Ders Adı', 'Ders Adı (EN)', 'T+U+L', 'Kredi', 'AKTS', 'Tür'], rows, [9, 12, 52, 52, 8, 7, 7, 10], first=True)
    rows = []
    for s in MUF['semesters']:
        for e in s['electiveSlots']:
            opts = 'Üniversite seçmelileri' if e['options'] == 'universityElectives' else ', '.join(e['options'])
            rows.append([s['semester'], e['slot'], e['type'], e['name'], e['nameEn'], e['ects'], opts])
    sheet(wb, 'Seçmeli Haklar', ['Yarıyıl', 'Hak', 'Tür', 'Adı', 'Adı (EN)', 'AKTS', 'Seçenekler'], rows, [9, 10, 7, 34, 34, 7, 120])
    rows = []
    for u in MUF['universityElectives']['courses']:
        for sl in u['slots']:
            rows.append([u['code'], u['name'], u.get('nameEn', ''), ['Pazartesi', 'Salı', 'Çarşamba', 'Perşembe', 'Cuma', 'Cumartesi', 'Pazar'][sl['day']], sl['start'], sl['end'], 'Online'])
    sheet(wb, 'Üniversite Seçmelileri', ['Ders Kodu', 'Ders Adı', 'Ders Adı (EN)', 'Gün', 'Başlangıç', 'Bitiş', 'Derslik'], rows, [12, 30, 30, 12, 10, 10, 10])
    rows = [[c, p] for c, ps in MUF['prerequisites'].items() for p in ps]
    sheet(wb, 'Ön Koşullar', ['Ders Kodu', 'Ön Koşul'], rows, [12, 12])
    rows = [['SE 1115', 'ENGR 1115'], ['MATH 1101', 'MATH 1100'], ['ME 1222', 'ME 2210'], ['CHEM 1120', 'CHEM 1130']]
    sheet(wb, 'Eşdeğerlikler', ['Eski Ders Kodu', 'Müfredattaki Karşılığı'], rows, [16, 22])
    wb.save(OUT / 'mufredat.xlsx')


# ---------- 2. Ders programları (bölümün dosyaları) ----------
def copy_timetables():
    for f in (ROOT / 'data').glob('*haftalık ders programı.xlsx'):
        y = f.name[0]
        tm = 'guz' if 'güz' in f.name else 'bahar'
        shutil.copy(f, OUT / f'ders-programi-{y}-sinif-{tm}.xlsx')


# ---------- 3. Kontenjan ----------
def make_quota():
    prog = json.loads((ROOT / 'data/program.json').read_text())
    full = {('ESE 3260', 1), ('COMP 4446', 1), ('ENGR 4400', 2), ('ME 4340', 1)}
    rows, seen = [], set()
    for y in ['3', '4']:
        for s in prog['terms']['guz'][y]:
            k = (s['code'], s['section'])
            if k in seen:
                continue
            seen.add(k)
            cap = 8 if s['code'] == 'ME 4910' else 40
            reg = cap if k in full else random.randint(cap // 3, cap - 2)
            rows.append(['2026-2027 Güz', s['code'], s['section'], cap, reg])
    wb = openpyxl.Workbook()
    sheet(wb, 'Kontenjan', ['Dönem', 'Ders Kodu', 'Şube', 'Kontenjan', 'Kayıtlı'], sorted(rows, key=lambda r: (r[1], r[2])), [16, 12, 7, 11, 9], first=True)
    wb.save(OUT / 'kontenjan.xlsx')


# ---------- 4. Akademik takvim ----------
CAL = [
    ('2026-08-31', '2026-09-04', 'Güz Yarıyılı Katkı Payı / Öğrenim Ücreti Ödemeleri'),
    ('2026-09-07', '2026-09-11', 'Güz Yarıyılı Ders Kayıtları ve Danışman Onayları'),
    ('2026-09-14', None, 'Güz Yarıyılı Derslerin Başlaması'),
    ('2026-09-21', '2026-09-25', 'Güz Yarıyılı Ders Ekle-Bırak'),
    ('2026-10-28', '2026-10-29', 'Cumhuriyet Bayramı (Resmi Tatil)'),
    ('2026-11-09', '2026-11-20', 'Güz Yarıyılı Ara Sınavlar'),
    ('2026-11-27', None, 'Dersten Çekilme için Son Gün'),
    ('2026-12-25', None, 'Güz Yarıyılı Derslerin Sona Ermesi'),
    ('2027-01-01', None, 'Yılbaşı (Resmi Tatil)'),
    ('2027-01-04', '2027-01-16', 'Güz Yarıyılı Final Sınavları'),
    ('2027-01-21', None, 'Final Sınavı Notlarının İlanı'),
    ('2027-01-25', '2027-01-29', 'Güz Yarıyılı Bütünleme Sınavları'),
    ('2027-02-08', '2027-02-12', 'Bahar Yarıyılı Ders Kayıtları ve Danışman Onayları'),
    ('2027-02-15', None, 'Bahar Yarıyılı Derslerin Başlaması'),
    ('2027-02-22', '2027-02-26', 'Bahar Yarıyılı Ders Ekle-Bırak'),
    ('2027-04-05', '2027-04-16', 'Bahar Yarıyılı Ara Sınavlar'),
    ('2027-05-19', None, "Atatürk'ü Anma, Gençlik ve Spor Bayramı (Resmi Tatil)"),
    ('2027-05-28', None, 'Bahar Yarıyılı Derslerin Sona Ermesi'),
    ('2027-06-01', '2027-06-12', 'Bahar Yarıyılı Final Sınavları'),
]


def make_calendar():
    wb = openpyxl.Workbook()
    d = lambda s: dt.datetime.strptime(s, '%Y-%m-%d') if s else None
    ws = sheet(wb, 'Akademik Takvim', ['Başlangıç', 'Bitiş', 'Etkinlik'], [[d(a), d(b), e] for a, b, e in CAL], [13, 13, 60], first=True)
    for row in ws.iter_rows(min_row=2, max_col=2):
        for c in row:
            c.number_format = 'DD.MM.YYYY'
    ws.insert_rows(1)
    ws['A1'] = 'ÖRNEK 2026-2027 Akademik Takvimi (demo amaçlı, resmi değildir)'
    ws['A1'].font = Font(bold=True, size=13)
    wb.save(OUT / 'akademik-takvim.xlsx')


# ---------- 5. Sınav takvimi ----------
def make_exams():
    plan = [
        ('ENGR 4400', 'Ara Sınav', '2026-11-10', '10:00', '11:00', 'X001'),
        ('ENGR 4411', 'Ara Sınav', '2026-11-11', '13:30', '15:00', 'X001'),
        ('ME 3220', 'Ara Sınav', '2026-11-12', '09:00', '11:00', 'M528'),
        ('ME 4240', 'Ara Sınav', '2026-11-12', '13:00', '15:00', 'M527'),
        ('ME 4330', 'Ara Sınav', '2026-11-16', '10:00', '12:00', 'M531'),
        ('EEE 4122', 'Ara Sınav', '2026-11-17', '14:00', '16:00', 'F004'),
        ('ME 4610', 'Ara Sınav', '2026-11-18', '10:00', '12:00', 'M531'),
        ('ESE 4240', 'Ara Sınav', '2026-11-19', '13:00', '15:00', 'M531'),
        ('ME 4910', 'Ara Sunum', '2026-11-20', '09:00', '17:00', 'M534'),
        ('ENGR 4411', 'Final', '2027-01-05', '13:30', '15:30', 'X001'),
        ('ME 3220', 'Final', '2027-01-06', '09:00', '11:00', 'M528'),
        ('ME 4240', 'Final', '2027-01-06', '14:00', '16:00', 'M527'),
        ('ME 4330', 'Final', '2027-01-08', '10:00', '12:00', 'M531'),
        ('EEE 4122', 'Final', '2027-01-11', '14:00', '16:00', 'F004'),
        ('ME 4610', 'Final', '2027-01-12', '10:00', '12:00', 'M531'),
        ('ESE 4240', 'Final', '2027-01-13', '13:00', '15:00', 'M531'),
        ('ENGR 4400', 'Final', '2027-01-14', '10:00', '11:00', 'X001'),
        ('ME 4910', 'Final Sunumu', '2027-01-15', '09:00', '17:00', 'M534'),
    ]
    rows = [[c, COURSES.get(c, {}).get('name', ''), k, dt.datetime.strptime(d, '%Y-%m-%d'), a, b, r] for c, k, d, a, b, r in plan]
    wb = openpyxl.Workbook()
    ws = sheet(wb, 'Sınav Takvimi', ['Ders Kodu', 'Ders Adı', 'Sınav Türü', 'Tarih', 'Başlangıç', 'Bitiş', 'Derslik'], rows, [12, 44, 13, 12, 10, 10, 10], first=True)
    for c in ws['D'][1:]:
        c.number_format = 'DD.MM.YYYY'
    wb.save(OUT / 'sinav-takvimi.xlsx')


# ---------- 6. Transkript (kurgusal Demo Öğrenci) ----------
TERMS = [
    ('2023-2024 Güz', [('ENGR 1115', 'B-'), ('MATH 1100', 'C'), ('MATH 1131', 'C+'), ('PHYS 1121', 'C-'), ('SOFL 1101', 'B+'), ('UMFD 1020', 'A-')]),
    ('2023-2024 Bahar', [('CHEM 1130', 'F'), ('MATH 1132', 'C'), ('ME 1110', 'B'), ('PHYS 1122', 'D+'), ('SOFL 1102', 'B'), ('UMFD 1120', 'A'), ('UMFD 1220', 'A-')]),
    ('2024-2025 Güz', [('EEE 2280', 'D'), ('ENGR 2120', 'A-'), ('HIST 1110', 'A-'), ('ISG 9110', 'A'), ('ISG 9210', 'A-'), ('ME 2210', 'C'), ('ME 2710', 'B+'), ('TURK 1110', 'A')]),
    ('2024-2025 Bahar', [('CHEM 1130', 'B-'), ('HIST 1210', 'A'), ('MATH 2230', 'B'), ('MATH 2263', 'C-'), ('ME 2220', 'D+'), ('ME 2320', 'C'), ('TURK 1210', 'B+'), ('UFND 7010', 'S')]),
    ('2025-2026 Güz', [('EEE 3718', 'B'), ('ENGR 3450', 'B-'), ('ME 3220', 'F'), ('ME 3320', 'C+'), ('ME 3420', 'B-'), ('UNV 1003', 'A')]),
    ('2025-2026 Bahar', [('MATH 2260', 'B+'), ('ME 3330', 'C-'), ('ME 3410', 'C'), ('ME 3510', 'B-'), ('ME 4230', 'B')]),
]
RANGE = {'A': (90, 98), 'A-': (85, 89), 'B+': (80, 84), 'B': (75, 79), 'B-': (70, 74), 'C+': (65, 69), 'C': (60, 64), 'C-': (55, 59), 'D+': (50, 54), 'D': (45, 49), 'F': (18, 40), 'S': (70, 95)}
UNV = {u['code']: u for u in MUF['universityElectives']['courses']}


def course_info(code):
    if code in COURSES:
        c = COURSES[code]
        return c['name'], c.get('credit', 0), c['ects'], 'Z' if any(code in s['mandatory'] for s in MUF['semesters']) else 'S'
    u = UNV[code]
    return u['name'].upper(), 2, 4, 'S'


def transcript_rows():
    rows = []
    for term, lst in TERMS:
        for code, letter in lst:
            nm, cr, ects, zs = course_info(code)
            lo, hi = RANGE[letter]
            rows.append((term, code, nm, cr, ects, random.randint(lo, hi), letter, zs))
    return rows


def make_transcript():
    rows = transcript_rows()
    wb = openpyxl.Workbook()
    sheet(wb, 'Transkript', ['Dönem', 'Ders Kodu', 'Ders Adı', 'Kredi', 'AKTS', 'Başarı Notu', 'Harf Notu', 'Z/S'], [list(r) for r in rows], [17, 12, 48, 7, 7, 11, 10, 6], first=True)
    sheet(wb, 'Öğrenci', ['Alan', 'Değer'], [['Adı Soyadı', 'DEMO ÖĞRENCİ'], ['Öğrenci No', '23070001234'], ['Fakülte', 'MÜHENDİSLİK FAKÜLTESİ'], ['Program', 'MAKİNE MÜHENDİSLİĞİ'], ['Kayıt Tarihi', '29.08.2023']], [16, 30])
    wb.save(OUT / 'transkript.xlsx')

    # PDF: OBS transkriptine benzer iki sütunlu düzen
    W, H = landscape(A4)
    c = canvas.Canvas(str(OUT / 'transkript.pdf'), pagesize=(W, H))
    c.setTitle('Transkript - Demo Öğrenci (örnek)')

    def header(y):
        c.setFont('DVB', 13); c.drawString(30, y, 'ÖRNEK TRANSKRİPT (demo amaçlı, resmi belge değildir)')
        c.setFont('DV', 9)
        c.drawString(30, y - 18, 'Adı Soyadı: DEMO ÖĞRENCİ      Öğrenci No: 23070001234      Fakülte: MÜHENDİSLİK FAKÜLTESİ      Program: MAKİNE MÜHENDİSLİĞİ      Kayıt Tarihi: 29.08.2023')
        return y - 44

    y = header(H - 36)
    colx = [30, W / 2 + 6]
    colw = W / 2 - 36
    by_term = {}
    for r in rows:
        by_term.setdefault(r[0], []).append(r)
    terms = list(by_term)
    for i in range(0, len(terms), 2):
        pair = terms[i:i + 2]
        need = 34 + 14 * max(len(by_term[t]) for t in pair) + 22
        if y - need < 30:
            c.showPage(); y = header(H - 36)
        for j, term in enumerate(pair):
            x = colx[j]
            yy = y
            c.setFillColorRGB(0.92, 0.94, 0.97); c.rect(x - 4, yy - 4, colw, 16, stroke=0, fill=1); c.setFillColorRGB(0, 0, 0)
            c.setFont('DVB', 9); c.drawString(x, yy, term)
            yy -= 16
            c.setFont('DVB', 7.5)
            for label, dx in [('Ders Kodu', 0), ('Ders Adı', 56), ('Kredi', colw - 150), ('AKTS', colw - 120), ('Ort', colw - 90), ('Harf', colw - 62), ('Z', colw - 30)]:
                c.drawString(x + dx, yy, label)
            yy -= 13
            c.setFont('DV', 7.5)
            for (_, code, nm, cr, ects, sc, letter, zs) in by_term[term]:
                nm2 = nm if len(nm) <= 36 else nm[:35] + '…'
                for txt, dx in [(code, 0), (nm2, 56), (str(cr), colw - 150), (str(ects), colw - 120), (str(sc), colw - 90), (letter, colw - 62), (zs, colw - 30)]:
                    c.drawString(x + dx, yy, txt)
                yy -= 14
        y -= need
    c.save()


# ---------- 7. Örnek yönetmelik ----------
REG = [
    ('ÖRNEK ÜNİVERSİTESİ ÖNLİSANS VE LİSANS EĞİTİM-ÖĞRETİM YÖNETMELİĞİ', 'DVB', 12),
    ('(Demo amaçlı hazırlanmış kısaltılmış örnek metindir; resmi yönetmelik değildir.)', 'DV', 9),
    ('', 'DV', 9),
    ('MADDE 8 – Öğretim süresi', 'DVB', 10),
    ('(1) Lisans programlarının normal öğretim süresi 4 yıl, azami öğretim süresi 7 yıldır.', 'DV', 10),
    ('', 'DV', 9),
    ('MADDE 14 – Ders yükü', 'DVB', 10),
    ('(1) Öğrenci her yarıyılda en fazla 30 AKTS ders alır. Genel not ortalaması 2,00 ve üzeri olan öğrenciler', 'DV', 10),
    ('danışmanının onayı ile yarıyıl başına en fazla 40 AKTS, ortalaması 2,00 altında olanlar en fazla 32 AKTS alabilir.', 'DV', 10),
    ('', 'DV', 9),
    ('MADDE 20 – Notlar', 'DVB', 10),
    ('(1) Başarı notları ve katsayıları şunlardır: A 4,00; A- 3,70; B+ 3,30; B 3,00; B- 2,70; C+ 2,30; C 2,00;', 'DV', 10),
    ('C- 1,70; D+ 1,30; D 1,00; F 0,00. S (yeterli) notu ortalamaya katılmaz.', 'DV', 10),
    ('(2) Ortalamalar derslerin AKTS değerleri ile ağırlıklandırılarak hesaplanır. Tekrar edilen derslerde son alınan not geçerlidir.', 'DV', 10),
    ('', 'DV', 9),
    ('MADDE 22 – Ders tekrarı', 'DVB', 10),
    ('(1) F notu alınan dersler tekrar edilir. Not yükseltmek amacıyla C- ve daha düşük harf notu alınan dersler tekrar alınabilir.', 'DV', 10),
    ('', 'DV', 9),
    ('MADDE 30 – Mezuniyet', 'DVB', 10),
    ('(1) Öğrencinin mezun olabilmesi için kayıtlı olduğu programın müfredatındaki bütün zorunlu ve seçmeli dersleri', 'DV', 10),
    ('başarıyla tamamlaması, en az 240 AKTS kazanması ve genel not ortalamasının en az 2,00 olması gerekir.', 'DV', 10),
    ('(2) Öğrenci 20 iş günü zorunlu stajını tamamlamadan mezun olamaz.', 'DV', 10),
]


def make_regulation():
    c = canvas.Canvas(str(OUT / 'yonetmelik-ornek.pdf'), pagesize=A4)
    c.setTitle('Örnek Yönetmelik (demo)')
    y = A4[1] - 60
    from reportlab.lib.utils import simpleSplit
    for txt, f, sz in REG:
        for line in (simpleSplit(txt, f, sz, A4[0] - 100) or ['']):
            c.setFont(f, sz); c.drawString(50, y, line); y -= sz + 6
        y -= 2
    c.save()


def make_regulation_txt():
    (OUT / 'yonetmelik-ornek.txt').write_text('\n'.join(t for t, f, s in REG), encoding='utf-8')


make_regulation_txt(); make_curriculum(); copy_timetables(); make_quota(); make_calendar(); make_exams(); make_transcript(); make_regulation()
for f in sorted(OUT.iterdir()):
    print(f.name, f.stat().st_size)
